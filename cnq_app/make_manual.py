"""Generate ``MANUAL.docx`` -- the end-user manual -- with python-docx.

Run from ``cnq_app/`` (needs the optional docs dependency):

    pip install python-docx
    python make_manual.py            # writes MANUAL.docx
    python make_manual.py -o out.docx

The manual is kept as this script (version-controlled, reproducible); the built
``MANUAL.docx`` is committed alongside it for convenience.
"""
from __future__ import annotations

import argparse
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = HERE / "MANUAL.docx"


def build(path: Path) -> Path:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as exc:  # noqa: BLE001
        raise SystemExit("python-docx is required: pip install python-docx") from exc

    doc = Document()
    normal = doc.styles["Normal"].font
    normal.name = "Calibri"
    normal.size = Pt(11)

    # ---- helpers ---------------------------------------------------------- #
    def h1(t):
        doc.add_heading(t, level=1)

    def h2(t):
        doc.add_heading(t, level=2)

    def para(t="", *, bold=False, italic=False):
        p = doc.add_paragraph()
        run = p.add_run(t)
        run.bold, run.italic = bold, italic
        return p

    def bullets(items):
        for it in items:
            doc.add_paragraph(str(it), style="List Bullet")

    def steps(items):
        for it in items:
            doc.add_paragraph(str(it), style="List Number")

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, head in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(head)
            run.bold = True
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)
        doc.add_paragraph()

    # ---- title ------------------------------------------------------------ #
    title = doc.add_heading("CNQ — User Manual", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = para("Censored Non-crossing Quantile Regression — a local app for "
               "predicting the full survival-time distribution.")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para("Everything runs on your computer; your data never leaves it. This manual "
         "covers training, auto-tuning, saving, predicting and projecting. It is "
         "generated from make_manual.py.").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # ---- 1. what it does -------------------------------------------------- #
    h1("1. What CNQ does")
    para("CNQ predicts, for each subject, the whole distribution of survival time — "
         "not just a single number. It does this by predicting a set of quantiles of "
         "the survival time (e.g. the 10th, 50th and 90th percentiles), from which you "
         "get a median prediction and an interval. The models are 'non-crossing', so "
         "the predicted quantiles never overlap.")
    para("Three things you can do:")
    bullets([
        "Train a model on your survival data (or let Auto-tune choose one).",
        "Predict survival-time quantiles for new subjects from a saved model.",
        "Project population trends — how many events by a time, or the time until a "
        "target number of events.",
    ])
    para("It works with censored data (subjects who hadn't had the event by their "
         "last follow-up), which is the norm in survival analysis.")

    # ---- 2. getting started ---------------------------------------------- #
    h1("2. Getting started")
    steps([
        "Start the app (double-click run.bat / run.sh, or run `python run.py` in the "
        "cnq_app folder). The first launch installs what it needs.",
        "Open http://127.0.0.1:8000/ in your browser.",
        "If you just updated the app, press Ctrl+Shift+R (Cmd+Shift+R on Mac) to load "
        "the new version.",
    ])
    para("Keep the terminal window open while you use the app; closing it stops the "
         "server. Everything else is done in the browser — no typing commands.")
    para("Want to see it work first? Use the built-in demo — see section 9.", italic=True)

    # ---- 3. concepts ------------------------------------------------------ #
    h1("3. Key concepts")
    table(["Term", "What it means"], [
        ["Follow-up time", "Time from start until the event or the last time the subject "
         "was seen (a positive number)."],
        ["Event indicator", "1 = the event happened; 0 = censored (event-free at last "
         "contact)."],
        ["Censoring", "Not observing the event before follow-up ended. CNQ accounts for "
         "it (IPCW weighting)."],
        ["Quantile (τ)", "The τ-th percentile of survival time. τ=0.5 is the median; "
         "τ=0.9 is a late time by which 90% are predicted to have had the event."],
        ["Non-crossing", "Higher quantiles are always ≥ lower ones, so each subject's "
         "predicted survival curve is well-behaved."],
        ["Survival curve S(t)", "The probability of still being event-free at time t; "
         "S(t) = 1 − (fraction with the event by t)."],
    ])

    # ---- 4. data ---------------------------------------------------------- #
    h1("4. Preparing your data")
    para("Use one CSV, one row per subject, with a single header row and no blank "
         "rows. You need:")
    bullets([
        "a follow-up time column (positive numbers),",
        "an event column (1 = event, 0 = censored — or two other values you'll map),",
        "one or more numeric covariate (feature) columns,",
        "optionally a subject-ID column (a label, not used as a predictor).",
    ])
    para("Guidance: aim for at least ~100 events, and enough that each split keeps ≥10 "
         "events. The app checks this and warns you. Rows with blank/non-numeric cells "
         "in the columns you use are excluded (never guessed).")

    # ---- 5. train --------------------------------------------------------- #
    h1("5. Train a model")
    h2("5.1 Upload")
    para("On the Train tab, drag in your CSV (or click to choose it). No data yet? Use "
         "the demo data, the sample data, or download a template.")
    h2("5.2 Map columns")
    para("Choose which columns are the duration and event, optionally a subject-ID "
         "column and a time unit (e.g. 'months', shown on plots), and tick the "
         "covariates to include. If the event column isn't 0/1, you're asked which "
         "value means the event. A live panel shows errors (which block training) and "
         "warnings (which you acknowledge), plus a summary (rows used, events, "
         "censoring, duration range).")
    h2("5.3 Settings: Choose settings or Auto-tune")
    para("Choose settings (default): pick one or more models and their "
         "hyperparameters. 'Paper preset' reuses values tuned in the paper on a public "
         "dataset of similar size (the closest is marked 'suggested'); 'Custom' lets "
         "you set them yourself and is pre-filled from your data's size and covariate "
         "count.")
    para("Auto-tune: let the app search several models and settings on your data and "
         "pick the best.")
    steps([
        "Tick the model families to try — KAN-CNQ and MLP-CNQ by default. Transformers "
        "are opt-in and help only for nonstandard / cross-modality data.",
        "Choose a search effort: Quick (~6 trials), Standard (~12), or Thorough (~24, "
        "with repeated splits for steadier picks).",
        "Click 'Find best settings' and watch the progress (best-so-far is shown). You "
        "can cancel; you'll still get the best found so far.",
        "Review the leaderboard (lower validation pinball is better), then click 'Use "
        "best settings' to fill in the winner. Review and Start training.",
    ])
    h2("5.4 Quantile levels")
    para("The default (0.1, 0.25, 0.5, 0.75, 0.9) is a good start. 'Every 10% / 5% / "
         "1%' or a custom list give a smoother, further-reaching predicted curve — "
         "useful for survival curves and for projection. 0.1, 0.5 and 0.9 are always "
         "included. Very low/high levels rely on few events and can be unreliable.")
    h2("5.5 Splits, seed, deterministic")
    para("Set the train/validation/test split, the number of repeated splits (more "
         "gives error bars on the metrics), a random seed, and optionally Deterministic "
         "mode (slower but exactly reproducible). Then Start training and watch the "
         "progress bar; you can Cancel.")
    h2("5.6 Reading the results")
    para("The report shows metrics (mean ± SD across repeated splits) and plots:")
    table(["Metric", "How to read it"], [
        ["Pinball (IPCW)", "Overall accuracy of the predicted quantiles; lower is "
         "better."],
        ["ICP 80% / 50%", "Interval coverage — how often the true time fell in the "
         "80% / 50% predicted interval; closer to 80% / 50% is better."],
        ["Uno C-index", "Discrimination — ranking subjects by risk; 0.5 = chance, 1.0 "
         "= perfect."],
        ["Calibration", "Predicted vs observed coverage per τ; near the diagonal is "
         "good."],
        ["Crossing rate", "Should be 0% for the non-crossing (_gaps) models."],
    ])
    para("Plots include the Kaplan–Meier outcome curve, training curves, calibration, "
         "individual quantile profiles and (with a dense grid) survival curves, "
         "coverage by interval width, permutation importance, and pinball across "
         "splits. Open the report or download the results zip.")

    # ---- 6. save ---------------------------------------------------------- #
    h1("6. Save a model")
    para("On the results, use 'Save a model' to bundle a trained model as a "
         ".cnqmodel file for reuse in Predict / Project or to share.")
    table(["Bundle type", "What it saves"], [
        ["Final model: refit on all data", "Retrains the chosen model on every row "
         "(15% held out for early stopping). Usually the best choice."],
        ["Ensemble of the repeated-split models", "Keeps every split's model and "
         "averages them (needs >1 repeated split)."],
        ["Single split", "One split's model."],
    ])
    para("Only the most recent runs are kept on disk (the CNQ_KEEP_RUNS setting, "
         "default 20), so save any model you want to keep. A saved .cnqmodel is "
         "self-contained and isn't affected by that cleanup.")

    # ---- 7. predict ------------------------------------------------------- #
    h1("7. Predict new subjects")
    steps([
        "On the Predict tab, choose a saved model (its summary is shown) or upload a "
        ".cnqmodel file.",
        "Upload a CSV of the new subjects.",
        "Map columns: a table lists each model feature with its training range and a "
        "dropdown of your file's columns. Columns are auto-matched by name; it shows "
        "'N of M features matched'. If names differ, use 'Download template for this "
        "model', 'Match by position' (confirmed), or pick columns yourself (a likely "
        "one is marked 'suggested'). Optionally set a subject-ID column, and — only if "
        "your subjects have known outcomes — time and event columns for external "
        "validation.",
        "Run. You get a predictions table and downloads: predictions CSV, an HTML "
        "report with plots, and a zip.",
    ])
    para("The predictions CSV has, per subject: each quantile q_<τ>, the median, the "
         "80% interval (low/high/width) and an out_of_range flag (1 if any feature is "
         "outside the training range — an extrapolation).")
    para("External validation (optional): if you provided time and event columns, the "
         "report adds IPCW pinball, 50%/80% coverage, Uno C-index, per-τ calibration "
         "and an observed-vs-predicted survival plot, using a censoring estimate from "
         "the new data. Only meaningful on subjects the model was NOT trained on.")

    # ---- 8. project ------------------------------------------------------- #
    h1("8. Project population trends")
    para("The Project tab answers aggregate questions instead of per-subject ones:")
    bullets([
        "Events by a time: enter time points → expected cumulative events (with a "
        "band).",
        "Time for a target: enter a number of events → the time by which they're "
        "expected (with a band).",
    ])
    para("Pick a model, then choose a basis:")
    table(["Basis", "What it uses"], [
        ["Population size N (headline)", "Scales the model's training-population "
         "Kaplan–Meier curve to N subjects. No cohort upload needed."],
        ["Upload a cohort", "Uses the model on those subjects' covariates (same "
         "column-mapping table as Predict); adjusts the trend for a cohort whose mix "
         "differs from training."],
    ])
    para("You get a cumulative-events-over-time chart with a band, answers to your "
         "questions, and downloads. Projection is only shown up to the observed "
         "follow-up horizon (the last event time in training); a time — or a target — "
         "beyond that is flagged, not extrapolated. A dense quantile grid gives a "
         "smoother, further-reaching curve.")

    # ---- 9. demo ---------------------------------------------------------- #
    h1("9. The built-in demo")
    para("Everything can be tried with simulated demo data, entirely in the browser:")
    steps([
        "Train tab → 'Use demo training data' → pick a model → Every 5% quantiles → "
        "Start training.",
        "Predict tab → 'Demo model (simulated data)'. If it says 'not built yet', click "
        "'Build demo model' once (about 2 minutes) — the data and model are created "
        "automatically.",
        "'Or use demo new subjects' → 20 new patients → Run (note the two flagged "
        "out-of-range). Then 500 with outcomes (coverage near nominal) and 500 shifted "
        "population (calibration worsens — why you validate on your own data).",
        "Project tab → pick the demo model → Population size N → enter time points or a "
        "target number of events.",
    ])

    # ---- 10. interpreting / caveats -------------------------------------- #
    h1("10. Interpreting predictions — and their limits")
    bullets([
        "Between-subject spread, not model uncertainty: a subject's quantile interval "
        "describes how outcomes vary between similar subjects, not the model's own "
        "uncertainty about its parameters. Projection bands are aggregate "
        "outcome/sampling uncertainty (Poisson-binomial for a cohort, Greenwood for "
        "the population).",
        "Population shift: predictions assume new subjects resemble the training "
        "population. A different hospital, era or inclusion criteria can make all "
        "predictions off — even in-range ones.",
        "Out-of-range flags: a subject with a feature outside the training range is "
        "flagged; those predictions are extrapolations and less reliable.",
        "External validation needs unseen data: outcome-based metrics are only "
        "meaningful on subjects the model was not trained on.",
        "Projection follow-up limit: trends are supported only up to the last observed "
        "event time; beyond it (or for targets needing more events than occur within "
        "follow-up) results aren't shown — extrapolating the censored tail would need "
        "a parametric model.",
    ])

    # ---- 11. troubleshooting --------------------------------------------- #
    h1("11. Troubleshooting")
    table(["Problem", "Fix"], [
        ["'Port 8000 is already in use'", "Run `python run.py 8001` and open "
         ":8001 instead."],
        ["Red banner about missing packages", "Check your internet, then "
         "`python run.py --reinstall`."],
        ["Banner: deepcnq repository not found", "Put the deepcnq checkout next to the "
         "app (../deepcnq) or set CNQ_REPO; then reload."],
        ["'No Python at …'", "Delete the .venv folder in cnq_app and run again — it "
         "rebuilds the environment."],
        ["Page looks plain / buttons do nothing", "Use exactly http://127.0.0.1:8000/ "
         "and hard-refresh (Ctrl+Shift+R)."],
        ["Demo model 'not built yet'", "Click 'Build demo model' in the Predict tab "
         "(about 2 minutes)."],
        ["Population projection unavailable", "The model has no stored training curve; "
         "re-save (or rebuild) the model to enable it."],
        ["App inside OneDrive/iCloud/Dropbox misbehaves", "Move it to a local, "
         "non-synced folder; cloud sync can lock the .venv."],
    ])

    # ---- 12. glossary ----------------------------------------------------- #
    h1("12. Glossary")
    table(["Term", "Meaning"], [
        ["IPCW", "Inverse-probability-of-censoring weighting — corrects metrics for "
         "censoring."],
        ["Pinball loss", "The quantile-regression loss; lower means better-calibrated, "
         "sharper quantiles."],
        ["Kaplan–Meier (KM)", "Non-parametric estimate of the survival curve from "
         "observed times/events."],
        ["Greenwood variance", "Standard variance formula for the KM estimate; gives "
         "the population projection band."],
        ["Uno C-index", "Censoring-adjusted concordance (ranking) measure."],
        ["ICP", "Interval coverage probability — how often the truth falls in a "
         "predicted interval."],
        ["Horizon", "The last observed event time; projection is supported up to here."],
        [".cnqmodel", "A self-contained saved model bundle (weights + metadata)."],
    ])

    doc.save(str(path))
    return path


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Build the CNQ user manual (MANUAL.docx).")
    parser.add_argument("-o", "--out", default=str(OUT), help="output .docx path")
    args = parser.parse_args(argv)
    out = build(Path(args.out))
    print(f"Wrote {out} ({out.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
